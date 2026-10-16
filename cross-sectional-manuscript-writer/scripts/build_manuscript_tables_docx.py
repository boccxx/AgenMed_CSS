#!/usr/bin/env python3
from __future__ import annotations

import argparse
import csv
import json
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


def read_json(path: Path):
    if not path.exists():
        raise SystemExit(f"Required registry not found: {path}")
    return json.loads(path.read_text(encoding="utf-8"))


def read_csv(path: Path) -> tuple[list[str], list[list[str]]]:
    if not path.exists():
        raise SystemExit(f"Registered table source does not exist: {path}")
    with path.open(newline="", encoding="utf-8-sig") as fh:
        reader = csv.DictReader(fh)
        headers = list(reader.fieldnames or [])
        rows = [[str(row.get(header, "")) for header in headers] for row in reader]
    if not headers:
        raise SystemExit(f"Registered table has no header: {path}")
    return headers, rows


def humanize(value: str) -> str:
    return value.replace("_", " ").strip().title()


def set_defaults(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    for style_name in ("Normal", "Heading 1"):
        doc.styles[style_name].font.name = "Times New Roman"
        doc.styles[style_name].font.size = Pt(11 if style_name == "Normal" else 12)


def set_cell_text(cell, text: str, *, bold: bool = False) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.name = "Times New Roman"
    run.font.size = Pt(9)


def set_border(cell, edge_name: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    edge = borders.find(qn(f"w:{edge_name}"))
    if edge is None:
        edge = OxmlElement(f"w:{edge_name}")
        borders.append(edge)
    edge.set(qn("w:val"), "single")
    edge.set(qn("w:sz"), "8")
    edge.set(qn("w:color"), "000000")


def clear_borders(table) -> None:
    borders = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        edge = OxmlElement(f"w:{side}")
        edge.set(qn("w:val"), "nil")
        borders.append(edge)
    table._tbl.tblPr.append(borders)


def add_registry_table(doc: Document, number: int, record: dict) -> None:
    source = Path(str(record.get("source_path", "")))
    headers, rows = read_csv(source)
    caption = record.get("caption") or record.get("display_title") or record.get("title") or record.get("table_id")
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(f"Table {number}. {caption}")
    run.bold = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(10)

    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    clear_borders(table)
    for idx, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[idx], humanize(header), bold=True)
        set_border(table.rows[0].cells[idx], "top")
        set_border(table.rows[0].cells[idx], "bottom")
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            set_cell_text(cells[idx], value)
    for cell in table.rows[-1].cells:
        set_border(cell, "bottom")

    note = record.get("note") or "Values are reproduced from the frozen registered source table."
    note_paragraph = doc.add_paragraph()
    note_run = note_paragraph.add_run(f"Note: {note}")
    note_run.italic = True
    note_run.font.name = "Times New Roman"
    note_run.font.size = Pt(9)
    doc.add_paragraph(f"Registry ID: {record.get('table_id', '')}")


def main() -> int:
    parser = argparse.ArgumentParser(description="Build manuscript tables from table_registry.json.")
    parser.add_argument("project_root")
    parser.add_argument("--output", required=True)
    parser.add_argument("--include-supplement", action="store_true")
    args = parser.parse_args()

    root = Path(args.project_root)
    records = read_json(root / "registries" / "table_registry.json")
    selected = [
        row for row in records
        if args.include_supplement or row.get("registry_role", "main_text") == "main_text"
    ]
    if not selected:
        raise SystemExit("No registered manuscript table is available.")

    doc = Document()
    set_defaults(doc)
    for number, record in enumerate(selected, start=1):
        add_registry_table(doc, number, record)
    output = Path(args.output)
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)
    print(f"Wrote {len(selected)} registered table(s) to {output}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

